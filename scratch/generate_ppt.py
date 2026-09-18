import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
PURPLE_DARK = RGBColor(107, 33, 168)   # #6b21a8
PURPLE_PRIMARY = RGBColor(124, 58, 237) # #7c3aed
PURPLE_LIGHT = RGBColor(243, 232, 255) # #f3e8ff
TEXT_DARK = RGBColor(17, 24, 39)       # #111827
TEXT_MUTED = RGBColor(75, 85, 99)      # #4b5563
WHITE = RGBColor(255, 255, 255)
BG_LIGHT = RGBColor(249, 250, 251)    # #f9fafb
GREEN = RGBColor(16, 185, 129)
ORANGE = RGBColor(245, 158, 11)

def add_header(slide, title_text, category_text="CIVICCONNECT HACKATHON DECK"):
    # Header Banner
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = PURPLE_DARK
    top_bar.line.fill.background()

    tf = top_bar.text_frame
    tf.margin_left = Inches(0.8)
    tf.margin_top = Inches(0.15)
    
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(216, 180, 254)

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

blank_slide_layout = prs.slide_layouts[6]

# SLIDE 1: Title Slide
slide1 = prs.slides.add_slide(blank_slide_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = PURPLE_DARK
bg1.line.fill.background()

# Title text box
tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.5))
tf1 = tb1.text_frame
tf1.word_wrap = True

p_sub = tf1.paragraphs[0]
p_sub.text = "SMART CITY URBAN GOVERNANCE PLATFORM"
p_sub.font.size = Pt(14)
p_sub.font.bold = True
p_sub.font.color.rgb = RGBColor(216, 180, 254)
p_sub.alignment = PP_ALIGN.CENTER

p_main = tf1.add_paragraph()
p_main.text = "CivicConnect"
p_main.font.size = Pt(54)
p_main.font.bold = True
p_main.font.color.rgb = WHITE
p_main.alignment = PP_ALIGN.CENTER

p_desc = tf1.add_paragraph()
p_desc.text = "AI-Powered Civic Issue Reporting & Verified Departmental Resolution Engine"
p_desc.font.size = Pt(20)
p_desc.font.color.rgb = RGBColor(243, 232, 255)
p_desc.alignment = PP_ALIGN.CENTER

p_footer = tf1.add_paragraph()
p_footer.text = "\nHackathon Presentation Deck  |  38 Districts  |  7 Civic Departments"
p_footer.font.size = Pt(13)
p_footer.font.color.rgb = RGBColor(216, 180, 254)
p_footer.alignment = PP_ALIGN.CENTER


# SLIDE 2: Problem Statement & Innovation
slide2 = prs.slides.add_slide(blank_slide_layout)
add_header(slide2, "Problem Statement & Solution Vision")

# Problem Box
card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
card1.fill.solid()
card1.fill.fore_color.rgb = RGBColor(254, 242, 242)
card1.line.color.rgb = RGBColor(252, 165, 165)
tf = card1.text_frame
tf.margin_left = tf.margin_top = Inches(0.4)
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "⚠️ Existing Challenges in Municipal Systems"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = RGBColor(185, 28, 28)

problems = [
    "Unverified Resolutions: Workers close complaints without department head review.",
    "Lack of Citizen Visibility: No real-time tracking of repair progress.",
    "Fake/Spam Registrations: Lack of OTP authentication leads to invalid data.",
    "Manual Reporting Delay: Generating CSV/PDF records for department heads takes hours.",
    "Language & Accessibility Barriers: Non-technical users struggle with text forms."
]
for item in problems:
    p = tf.add_paragraph()
    p.text = "• " + item
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK

# Solution Box
card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.3))
card2.fill.solid()
card2.fill.fore_color.rgb = RGBColor(240, 253, 244)
card2.line.color.rgb = RGBColor(134, 239, 172)
tf = card2.text_frame
tf.margin_left = tf.margin_top = Inches(0.4)
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ CivicConnect Innovation & Solution"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = RGBColor(21, 128, 61)

solutions = [
    "Strict 5-Step Workflow: Complaints require Department Head approval before 'Resolved'.",
    "Live Progress Tracker: 5-stage visual tracker for citizens and admins.",
    "Brevo 2-Factor OTP: 6-digit OTP verification via Email & SMS with 5-min expiry.",
    "1-Click PDF Report Exporter: Custom scope reports (Date/Month/Year/Dept).",
    "AI Voice Complaints: Multi-lingual speech-to-text complaint logging."
]
for item in solutions:
    p = tf.add_paragraph()
    p.text = "• " + item
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK


# SLIDE 3: Multi-Role System Architecture
slide3 = prs.slides.add_slide(blank_slide_layout)
add_header(slide3, "Multi-Role Governance Architecture")

roles = [
    ("🧑‍💼 Citizen", "Files issue via text or AI Voice\nTracks live status\nReceives Email/SMS alerts\nViews full repair log", PURPLE_LIGHT, PURPLE_PRIMARY),
    ("👨‍🔧 Field Worker", "Accepts assigned tasks\nUpdates status to In Progress\nUploads completion photo & notes\nTriggers Review Stage", RGBColor(239, 246, 255), RGBColor(37, 99, 235)),
    ("🏢 Department Head", "Auto-receives district issues\nAssigns worker + completion deadline\nReviews worker work submission\nApproves & Marks Resolved", RGBColor(254, 243, 199), RGBColor(217, 119, 6)),
    ("👑 Super Admin", "Manages 38 districts & 7 depts\nMonitors master dashboard\nGenerates custom PDF reports\nManages Dept Heads & Workers", RGBColor(243, 244, 246), RGBColor(75, 85, 99))
]

for idx, (title, desc, bg_col, border_col) in enumerate(roles):
    x = Inches(0.8 + idx * 3.0)
    card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.7), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_col
    card.line.color.rgb = border_col
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.margin_left = tf.margin_top = Inches(0.3)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = border_col

    for line in desc.split('\n'):
        p = tf.add_paragraph()
        p.text = "✔ " + line
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK


# SLIDE 4: Strict 5-Step Complaint Verification Workflow
slide4 = prs.slides.add_slide(blank_slide_layout)
add_header(slide4, "Strict 5-Step Verification Workflow")

steps = [
    ("1. SUBMITTED", "Citizen logs civic issue\nAuto-routed to Dept Head", RGBColor(239, 246, 255), RGBColor(37, 99, 235)),
    ("2. ASSIGNED", "Dept Head sets deadline\nAssigns field worker", RGBColor(243, 232, 255), PURPLE_PRIMARY),
    ("3. IN PROGRESS", "Worker accepts task\nExecutes field repairs", RGBColor(254, 243, 199), RGBColor(217, 119, 6)),
    ("4. SENT FOR REVIEW", "Worker submits photo & notes\nPending DH Approval", RGBColor(254, 242, 242), RGBColor(220, 38, 38)),
    ("5. RESOLVED", "Dept Head approves work\nComplaint marked Resolved", RGBColor(240, 253, 244), RGBColor(21, 128, 61))
]

for idx, (title, desc, bg_col, border_col) in enumerate(steps):
    x = Inches(0.8 + idx * 2.4)
    box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.2), Inches(2.1), Inches(4.2))
    box.fill.solid()
    box.fill.fore_color.rgb = bg_col
    box.line.color.rgb = border_col
    box.line.width = Pt(2.5)

    tf = box.text_frame
    tf.margin_left = tf.margin_top = Inches(0.2)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = border_col

    for line in desc.split('\n'):
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK


# SLIDE 5: Security & Brevo OTP Authentication Engine
slide5 = prs.slides.add_slide(blank_slide_layout)
add_header(slide5, "Security & Brevo OTP Authentication Engine")

card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
card.fill.solid()
card.fill.fore_color.rgb = RGBColor(249, 250, 251)
card.line.color.rgb = PURPLE_PRIMARY

tf = card.text_frame
tf.margin_left = tf.margin_top = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🔒 Brevo Transactional Email & SMS OTP Infrastructure"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = PURPLE_DARK

features = [
    ("6-Digit Secure OTP", "Generated via java.security.SecureRandom (100000 - 999999)."),
    ("5-Minute Time Expiration", "Automated cache expiry; invalidates stale OTP codes after 300 seconds."),
    ("Max 5 Attempt Rate Limiter", "Protects against brute-force attacks; revokes OTP after 5 failed tries."),
    ("Brevo v3 REST API Integration", "Email via POST /v3/smtp/email and SMS via POST /v3/transactionalSMS/sms."),
    ("Zero Client API Leakage", "Brevo API Key stored strictly in Spring Boot backend environment variables."),
    ("JWT & Password Security", "BCrypt password hashing and Spring Security stateless JWT authorization.")
]

for title, desc in features:
    p = tf.add_paragraph()
    p.text = f"• {title}: "
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PURPLE_PRIMARY
    run = p.add_run()
    run.text = desc
    run.font.bold = False
    run.font.color.rgb = TEXT_DARK


# SLIDE 6: Automated PDF Report Exporter
slide6 = prs.slides.add_slide(blank_slide_layout)
add_header(slide6, "Automated PDF Report Generator")

rows, cols = 6, 4
left, top, width, height = Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2)
table_shape = slide6.shapes.add_table(rows, cols, left, top, width, height)
table = table_shape.table

headers = ["Admin Module", "Export Options Supported", "Included Data Fields", "Key Value"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = PURPLE_DARK
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE

data = [
    ("Citizens Directory", "All, By Date, By Month, By Year", "ID, Name, Email, Mobile, District, Address, Joining Date", "Complete citizen master audit"),
    ("Complaints Directory", "All, Filtered, By Date, By Month, By Year", "Complaint ID, Issue, Location, Priority, Status, Worker, Deadline", "Full complaint life cycle tracking"),
    ("Workers Directory", "All, By Date, By Month, By Year", "Worker ID, Name, Dept, District, Mobile/Email, Joining Date", "Field inspector tracking"),
    ("Dept Heads Directory", "All, By Date, By Month, By Year", "Name, Email, Mobile, Department, Joining Date, District", "Executive leadership directory"),
    ("Departments Directory", "All Departments, By Department", "Dept ID, Dept Code, Dept Name, Description, Total Officers", "Civic master department listing")
]

for row_idx, row_data in enumerate(data, start=1):
    for col_idx, cell_value in enumerate(row_data):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(249, 250, 251) if row_idx % 2 == 0 else WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = cell_value
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_DARK


# SLIDE 7: Technology Stack & Technical Rigor
slide7 = prs.slides.add_slide(blank_slide_layout)
add_header(slide7, "Technology Stack & Rigorous Engineering")

tech_boxes = [
    ("🎨 Frontend", "React 18  |  Vite\nTailwind CSS  |  Redux Toolkit\nLucide React Icons\njsPDF & AutoTable", PURPLE_LIGHT, PURPLE_PRIMARY),
    ("⚙️ Backend", "Spring Boot 3.2.4 (Java 21)\nSpring Data JPA & Hibernate\nSpring Security & JWT\nBrevo REST Client API", RGBColor(239, 246, 255), RGBColor(37, 99, 235)),
    ("🗄️ Database", "Oracle Database 21c XE\nOracle Thin Driver (OJDBC11)\nRelational JPA Mapping\nSequences & Foreign Keys", RGBColor(240, 253, 244), RGBColor(21, 128, 61))
]

for idx, (title, desc, bg_col, border_col) in enumerate(tech_boxes):
    x = Inches(0.8 + idx * 4.0)
    card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(3.7), Inches(5.0))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_col
    card.line.color.rgb = border_col
    card.line.width = Pt(2)

    tf = card.text_frame
    tf.margin_left = tf.margin_top = Inches(0.3)
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = border_col

    for line in desc.split('\n'):
        p = tf.add_paragraph()
        p.text = "• " + line
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK


# SLIDE 8: Hackathon Impact & Future Roadmap
slide8 = prs.slides.add_slide(blank_slide_layout)
add_header(slide8, "Hackathon Impact & Future Scope")

# Left Column - Impact
c1 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
c1.fill.solid()
c1.fill.fore_color.rgb = RGBColor(240, 253, 244)
c1.line.color.rgb = RGBColor(21, 128, 61)
tf = c1.text_frame
tf.margin_left = tf.margin_top = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🌟 Project Achievements & Impact"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = RGBColor(21, 128, 61)

impacts = [
    "Zero Unverified Resolutions: Mandatory Dept Head sign-off.",
    "Full State Coverage: Pre-loaded with 38 Tamil Nadu districts & 7 civic departments.",
    "Enterprise Security: 2-Factor Brevo OTP & stateless JWT authorization.",
    "Accessibility First: AI Voice complaint filing for non-technical citizens."
]
for item in impacts:
    p = tf.add_paragraph()
    p.text = "✔ " + item
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK

# Right Column - Roadmap
c2 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.6), Inches(5.2))
c2.fill.solid()
c2.fill.fore_color.rgb = PURPLE_LIGHT
c2.line.color.rgb = PURPLE_PRIMARY
tf = c2.text_frame
tf.margin_left = tf.margin_top = Inches(0.4)
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "🚀 Future Expansion Roadmap"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = PURPLE_DARK

roadmap = [
    "Interactive GIS Map: Live heatmap tracking of city infrastructure issues.",
    "Predictive AI Analytics: Forecasting recurring road & water damages.",
    "Native Mobile App: Offline-first Flutter app for field workers with auto-sync.",
    "IoT Integration: Automated sensor alerts for flooding and drainage overflow."
]
for item in roadmap:
    p = tf.add_paragraph()
    p.text = "📌 " + item
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_DARK


# Save PPTX to target directory
output_dir = r"C:\Users\Srimathi\.gemini\antigravity\brain\d6b435e6-b8ba-48bc-bca5-fd2ea3de1807"
os.makedirs(output_dir, exist_ok=True)
pptx_path = os.path.join(output_dir, "CivicConnect_Hackathon_Presentation.pptx")
prs.save(pptx_path)
print(f"Presentation saved successfully to {pptx_path}")
